"""Render stored report snapshots into downloadable PDF and Excel files."""

from __future__ import annotations

import os
import re
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.report import Report


EXPORT_DIRECTORY = Path(__file__).resolve().parents[2] / "storage" / "reports"


def _display(value: Any) -> str:
    """Produce a readable cell value without changing the stored snapshot."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_display(item)}" for key, item in value.items())
    if isinstance(value, list):
        return ", ".join(_display(item) for item in value)
    return str(value)


def _flatten_record(value: Any, prefix: str = "") -> dict[str, Any]:
    """Flatten nested snapshot dictionaries for a structured worksheet/table."""
    if not isinstance(value, dict):
        return {prefix or "value": value}
    flattened: dict[str, Any] = {}
    for key, item in value.items():
        column = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict):
            flattened.update(_flatten_record(item, column))
        else:
            flattened[column] = item
    return flattened


def _snapshot(report: Report) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
    """Validate that a generated report has a usable saved snapshot."""
    if report.status != "generated" or not isinstance(report.report_data, dict):
        raise ValueError("Report is not ready for export.")
    summary = report.report_data.get("summary", {})
    records = report.report_data.get("records", [])
    filters = report.report_data.get("filters", {})
    if not isinstance(summary, dict) or not isinstance(records, list) or not isinstance(filters, dict):
        raise ValueError("Report is not ready for export.")
    return summary, [item for item in records if isinstance(item, dict)], filters


def _safe_filename(report: Report, extension: str) -> str:
    report_type = re.sub(r"[^a-z0-9]+", "-", report.report_type.value.lower()).strip("-")
    return f"report-{report.id}-{report_type}.{extension}"


def _target_path(report: Report, extension: str) -> Path:
    EXPORT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    return EXPORT_DIRECTORY / _safe_filename(report, extension)


def _reference(path: Path) -> str:
    return path.relative_to(EXPORT_DIRECTORY.parents[1]).as_posix()


def _write_pdf(path: Path, report: Report, summary: dict[str, Any], records: list[dict[str, Any]], filters: dict[str, Any]) -> None:
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(
        str(path), pagesize=landscape(A4), rightMargin=1.2 * cm, leftMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm, pageCompression=0,
    )
    story = [Paragraph("Food Monitoring Report", styles["Title"])]
    metadata = [
        ["Report type", report.report_type.value],
        ["Generated date", _display(report.generated_at)],
        ["Date range", f"{_display(report.date_from)} to {_display(report.date_to)}" if report.date_from or report.date_to else "Not specified"],
        ["Record count", str(report.record_count)],
    ]
    story.extend([Spacer(1, 0.25 * cm), _table(metadata, [2.8 * cm, 21.5 * cm]), Spacer(1, 0.4 * cm)])
    story.append(Paragraph("Summary", styles["Heading2"]))
    summary_rows = [[str(key).replace("_", " ").title(), _display(value)] for key, value in summary.items()]
    story.extend([_table(summary_rows or [["Summary", "No summary values available"]], [5 * cm, 19.3 * cm]), Spacer(1, 0.4 * cm)])
    if filters:
        story.append(Paragraph("Filters", styles["Heading2"]))
        story.extend([_table([[str(key).replace("_", " ").title(), _display(value)] for key, value in filters.items()], [5 * cm, 19.3 * cm]), Spacer(1, 0.4 * cm)])
    story.append(Paragraph("Report Data", styles["Heading2"]))
    flattened = [_flatten_record(record) for record in records]
    columns = list(dict.fromkeys(key for record in flattened for key in record))
    if not columns:
        story.append(Paragraph("No records are available for this report.", styles["BodyText"]))
    else:
        data = [columns] + [[_display(record.get(column)) for column in columns] for record in flattened]
        width = 26.5 * cm / len(columns)
        story.append(_table(data, [width] * len(columns), header=True))
    document.build(story)


def _table(data: list[list[str]], widths: list[float], header: bool = False) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B7C4D1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")])
    table.setStyle(TableStyle(commands))
    return table


def _write_excel(path: Path, report: Report, summary: dict[str, Any], records: list[dict[str, Any]], filters: dict[str, Any]) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Report Summary"
    rows = [
        ("Report type", report.report_type.value),
        ("Generated date", _display(report.generated_at)),
        ("Date range", f"{_display(report.date_from)} to {_display(report.date_to)}" if report.date_from or report.date_to else "Not specified"),
        ("Record count", report.record_count),
    ]
    rows.extend((f"Summary: {key}", _display(value)) for key, value in summary.items())
    rows.extend((f"Filter: {key}", _display(value)) for key, value in filters.items())
    for key, value in rows:
        summary_sheet.append([key.replace("_", " ").title(), value])
    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 50
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)

    data_sheet = workbook.create_sheet("Report Data")
    flattened = [_flatten_record(record) for record in records]
    columns = list(dict.fromkeys(key for record in flattened for key in record))
    if columns:
        data_sheet.append(columns)
        for cell in data_sheet[1]:
            cell.font = Font(bold=True)
        for record in flattened:
            data_sheet.append([_display(record.get(column)) for column in columns])
        for index, column in enumerate(columns, 1):
            data_sheet.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 14), 35)
    else:
        data_sheet.append(["Report Data"])
        data_sheet.append(["No records are available for this report."])
        data_sheet["A1"].font = Font(bold=True)
        data_sheet.column_dimensions["A"].width = 45
    workbook.save(path)


def export_report(db: Session, report: Report, export_format: str) -> Path:
    """Create an export from a stored snapshot and save its safe relative reference."""
    summary, records, filters = _snapshot(report)
    extension = "pdf" if export_format == "pdf" else "xlsx"
    target = _target_path(report, extension)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=EXPORT_DIRECTORY, suffix=f".{extension}", delete=False) as handle:
            temporary = Path(handle.name)
        if export_format == "pdf":
            _write_pdf(temporary, report, summary, records, filters)
        else:
            _write_excel(temporary, report, summary, records, filters)
        os.replace(temporary, target)
        if export_format == "pdf":
            report.pdf_file_reference = _reference(target)
        else:
            report.excel_file_reference = _reference(target)
        db.commit()
        return target
    except Exception:
        db.rollback()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise
